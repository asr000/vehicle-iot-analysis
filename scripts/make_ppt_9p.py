"""按 9 页规划生成《车辆遥测数据分析汇报.pptx》。

口径与全部数字以 reports/审查与论证链条.md、reports/项目说明与汇报大纲.md 为准
（两份文档已经过逐数核对），notebook 仅用于必要时复核。
讲稿写入每页备注栏；不生成独立 doc 文件；不改动既有 汇报.pptx。
"""
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

ROOT = Path(__file__).resolve().parent.parent
FIG = ROOT / "figures"

DARK = RGBColor(0x1F, 0x3B, 0x5C)
GRAY = RGBColor(0x40, 0x40, 0x40)
LIGHT = RGBColor(0xEE, 0xF2, 0xF7)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]


def textbox(slide, left, top, width, height, lines, size=16, color=GRAY,
            bold=False, align=None):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line
        p.space_after = Pt(7)
        if align:
            p.alignment = align
        for run in p.runs:
            run.font.size = Pt(size)
            run.font.color.rgb = color
            run.font.name = "Microsoft YaHei"
            run.font.bold = bold
    return box


def title_bar(slide, text, page_no):
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, Inches(1.0))
    bar.fill.solid()
    bar.fill.fore_color.rgb = DARK
    bar.line.fill.background()
    tf = bar.text_frame
    tf.text = text
    tf.paragraphs[0].alignment = PP_ALIGN.LEFT
    tf.margin_left = Inches(0.5)
    for run in tf.paragraphs[0].runs:
        run.font.size = Pt(26)
        run.font.bold = True
        run.font.color.rgb = WHITE
        run.font.name = "Microsoft YaHei"
    textbox(slide, Inches(12.5), Inches(7.05), Inches(0.7), Inches(0.4),
            [str(page_no)], size=12)


def bullets(slide, items, left=Inches(0.5), top=Inches(1.3), width=Inches(6.7), size=16):
    return textbox(slide, left, top, width, Inches(5.5),
                   ["• " + t for t in items], size=size)


def picture(slide, name, left=Inches(7.4), top=Inches(1.3), width=Inches(5.5)):
    path = FIG / name
    if path.exists():
        slide.shapes.add_picture(str(path), left, top, width=width)


def notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


# ---------------- 第 1 页 封面 ----------------
s = prs.slides.add_slide(BLANK)
bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
bg.fill.solid(); bg.fill.fore_color.rgb = DARK; bg.line.fill.background()
textbox(s, Inches(1), Inches(2.6), Inches(11.3), Inches(1.2),
        ["车辆遥测数据分析与告警识别"], size=40, color=WHITE, bold=True)
textbox(s, Inches(1), Inches(3.9), Inches(11.3), Inches(1.8),
        ["数据清洗与规整 · 车辆运行状态分析 · 告警识别与成因判断",
         "汇报人：________　　课程：________"], size=18, color=LIGHT)
notes(s, "各位老师好，我汇报的题目是车辆物联网遥测数据分析与告警识别。数据来自三台联网车辆，"
         "约七天半，内容包括原始传感器的读数和告警标注。我的工作分四块：数据清洗、车辆运行特征分析、"
         "告警识别与成因判断、告警分类模型。下面按这个顺序讲，大约五分钟。")

# ---------------- 第 2 页 数据概况与预处理 ----------------
s = prs.slides.add_slide(BLANK)
title_bar(s, "从原始遥测数据到可分析的数据集", 2)
bullets(s, [
    "原始数据：173,853 行的长表，每行只有一条记录——某辆车、某时刻、某个传感器的一个读数",
    "清洗：统一时间格式；把挤在一个格子里的定位拆成纬度、经度、海拔三列；读数统一为数字",
    "规整：按“车辆 + 10 秒一格”拼成宽表，共 35,306 行 × 69 列；每格放这一格里的最后一次读数",
    "缺失处理：一格里没有读数就留空，不填 0、不编数据，缺多少后面都如实说明",
    "质量核对：清洗后 176,299 个数字逐个与原始数据比对，26 项检查全部通过",
])
steps = ["原始长表", "清洗与类型统一", "定位字段拆分", "时间对齐与长转宽", "分析数据集"]
x = Inches(0.6)
for st in steps:
    shp = s.shapes.add_shape(MSO_SHAPE.CHEVRON, x, Inches(6.0), Inches(2.5), Inches(0.7))
    shp.fill.solid(); shp.fill.fore_color.rgb = LIGHT; shp.line.color.rgb = DARK
    tf = shp.text_frame; tf.text = st
    for run in tf.paragraphs[0].runs:
        run.font.size = Pt(13); run.font.color.rgb = DARK; run.font.name = "Microsoft YaHei"
    x += Inches(2.45)
notes(s, "先说数据怎么处理的。原始数据十七万多行，但每行只记一个传感器的一个数，同一个时刻的车速、"
         "转速散在不同行里，没法直接比。所以做了三件事：统一时间格式；把定位字段一个格子里的三个数拆成三列；"
         "按每辆车每十秒一格拼成宽表，一行就是一辆车在某十秒内的完整状态。格子里没有读数就留空，不填零也不编数。"
         "处理完一共三万五千多行。质量上有独立程序把十七万六千多个数字逐个和原始数据比对，二十六项检查全部通过。")

# ---------------- 第 3 页 车辆运行状态 ----------------
s = prs.slides.add_slide(BLANK)
title_bar(s, "车辆运行状态的时间变化特征", 3)
bullets(s, [
    "车速和转速都不是连续曲线，而是“开一段、停很久、再开一段”",
    "转速形态与车速高度同步：踩油门转速升高、车速跟着升高",
    "车辆3 有一条约 860 转的稳定基线，对应怠速",
    "车辆2 出现过一次车速 175、转速 9000 的读数，与前后数据脱节，判定为设备误报，后续分析剔除",
    "注意：车速只有 6.4% 的格子有值，曲线空缺处是“没有上报”，不是 0",
], width=Inches(5.6))
picture(s, "B02_speed_timeseries.png", Inches(6.3), Inches(1.3), width=Inches(3.4))
picture(s, "B03_rpm_timeseries.png", Inches(9.9), Inches(1.3), width=Inches(3.4))
notes(s, "接下来看车是怎么开的。左边是车速、右边是转速随时间的曲线。两个共同点：一是数据是碎片化的，"
         "开一段、停很久、再开一段，曲线断开的地方是传感器没上报，不是车速为零；二是转速和车速高度同步，"
         "这符合开车的常识。车辆3 还有一条每分钟八百六十转的稳定基线，是怠速。另外车辆2 出现过一次车速一百七十五、"
         "转速九千的读数，和前后完全脱节，我们判定是设备误报，后面的分析把它剔除了。")

# ---------------- 第 4 页 三车对比 ----------------
s = prs.slides.add_slide(BLANK)
title_bar(s, "不同车辆的运行特征有何差异？", 4)
tbl = s.shapes.add_table(5, 4, Inches(0.6), Inches(1.5), Inches(7.6), Inches(3.4)).table
tbl.columns[0].width = Inches(3.4)
head = ["指标", "车辆1", "车辆2", "车辆3"]
data = [
    ["车速中位数 (km/h)", "49", "5", "22"],
    ["静止时间占比", "3.5%", "43.2%", "7.6%"],
    ["平均加速度绝对值 (m/s²)", "0.69", "1.14", "0.67"],
    ["异常读数", "无", "175 km/h 尖峰", "无"],
]
for j, h in enumerate(head):
    tbl.cell(0, j).text = h
for i, row in enumerate(data, 1):
    for j, v in enumerate(row):
        tbl.cell(i, j).text = v
for r in range(5):
    for c in range(4):
        for p in tbl.cell(r, c).text_frame.paragraphs:
            for run in p.runs:
                run.font.size = Pt(14); run.font.name = "Microsoft YaHei"
bullets(s, [
    "三台车工况差异明显：车辆1 快且路况好；车辆3 是城区低速；车辆2 基本不动",
    "车辆2 各项数值偏高，但可用样本只有三十几个，不能下结论",
    "数据量差很多：车辆3 占 93.4%，所以所有对比都分开三台车做",
], top=Inches(5.1), width=Inches(12.0))
picture(s, "B06_speed_bin_mix.png", Inches(8.5), Inches(1.4), width=Inches(4.3))
notes(s, "三台车对比。左表是关键指标：车辆1 车速中位数四十九，以中高速为主；车辆3 中位数二十二，城区低速；"
         "车辆2 中位数五，近一半时间不动，还有刚才说的异常尖峰。它的加速度数值也偏高，"
         "但可用样本只有三十几个，不能据此说这辆车开得猛。还要说明一点：三台车数据量差很多，车辆3 占了百分之九十三，"
         "所以我们所有对比都分车做，混在一起算平均会把差异盖住。")

# ---------------- 第 5 页 告警识别 ----------------
s = prs.slides.add_slide(BLANK)
title_bar(s, "如何识别告警事件？", 5)
bullets(s, [
    "识别依据：原始数据自带的告警标注 alarmClass（0–5），按 10 秒格子统计每类出现的次数",
    "对每一类告警分别做“出现/未出现”的判定，一个格子可以同时属于多类",
    "带告警的格子里，绝大多数同时带着好几类——只带一类的格子只有 14 个，所以不能按“一格一类”处理",
    "各类占比 7%–14%，都不算稀少；但按车差很多：4 号在车辆1 占 59.8%、在车辆3 只占 6.2%，5 号在车辆2 完全没有",
])
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
cd = CategoryChartData()
cd.categories = ["code1", "code2", "code3", "code4", "code5"]
cd.add_series("告警窗口数", (2822, 2521, 3164, 3155, 5031))
gf = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(7.6), Inches(1.6),
                        Inches(5.2), Inches(4.6), cd)
gf.chart.has_legend = False
gf.chart.chart_title.text_frame.text = "各类告警的窗口数"
notes(s, "怎么识别告警。依据是原始数据自带的告警标注，一共六类，零是正常，一到五是五种告警。"
         "我们按十秒格子统计每类出现的次数，对每一类分别判断它在这一格里出没出现——一个格子可以同时属于好几类，"
         "这是这份数据的特点：带告警的格子里绝大多数都同时带着好几类，只带一类的格子只有十四个，"
         "所以不能强行规定一格算哪一类。右图是各类的窗口数，占比百分之七到十四，都不算稀少。"
         "但按车看差很多，比如四号在车辆1 近六成的格子里出现，在车辆3 只有百分之六。")

# ---------------- 第 6 页 告警与状态关系 / 成因 ----------------
s = prs.slides.add_slide(BLANK)
title_bar(s, "告警发生时，车辆处于什么状态？", 6)
bullets(s, [
    "比法：在每一台车内部，比较“带该类告警的格子”和“正常格子”，避免车与车的差别混进来",
    "带 1–4 号告警的格子几乎全部处于点火状态（比正常格子高 0.68–0.86）；5 号没有这个现象",
    "带 1–4 号告警的格子，电池电压系统性偏高；5 号没有差别",
    "1–4 号告警常常连续出现几十秒到一两分钟；5 号几乎全是一次一闪而过",
    "可能的方向（只是统计关联，不是因果；五类编号对应的故障名称数据里没有，需数据提供方确认）："
    "1–4 号是“车开着、电气状态不对”时成组出现的一组告警；5 号是独立的另一类",
])
picture(s, "C04_effect_size.png", Inches(7.6), Inches(1.5), width=Inches(5.2))
notes(s, "这一页是成因分析，也是关键的一页。比法是在每台车内部比较带告警的格子和正常的格子，"
         "因为三台车本身差很多。右图是结果：左边蓝色是电池电压的差异，一到四号告警时电压系统性偏高，五号没有；"
         "右边是点火标志，一到四号告警几乎全部发生在车点着火的时候，五号没有关联。再补一个时间形态上的证据："
         "一到四号的告警常常连续出现一两分钟，五号几乎全是一次一闪而过。所以我们的判断是：一到四号像是"
         "车开着、电气状态不对时成组出现的一组告警，五号是独立的另一类。必须说明两点：这只是统计上的关联，"
         "数据分不出谁先谁后；而且五类编号对应的故障名称数据里没有提供，具体是什么故障要由数据提供方确认，我们没有猜测。")

# ---------------- 第 7 页 模型与评估 ----------------
s = prs.slides.add_slide(BLANK)
title_bar(s, "模型能否有效识别告警？", 7)
bullets(s, [
    "任务：给一个格子的 43 个传感器读数，自动判断五类告警各自是否出现；输出五个“是/否”",
    "模型：逻辑回归（各读数乘系数相加）与 XGBoost（一连串“转速是否大于某值”式的判断）；后者不怕格子空着",
    "考法：前 80% 的时间学、后 20% 的时间考，模拟用过去预测未来",
    "分数：五类各自算“查准、查全、F1”，再取平均（宏平均 F1）；不只看准确率——永远说“正常”的模型准确率 91.5%，"
    "但一个告警都没找到",
    "结果：宏平均 F1 从 0（全说正常）→ 0.60（逻辑回归）→ 0.78（XGBoost）；1–3 号 F1 0.80 以上，"
    "4 号 0.66、5 号 0.44，是两类难点",
], width=Inches(6.9))
picture(s, "C10_confusion.png", Inches(7.6), Inches(1.5), width=Inches(5.4))
notes(s, "模型部分。任务很直接：给一个格子的四十三个传感器读数，自动判断五类告警各自在不在。"
         "用了两个模型：逻辑回归，把读数各乘一个系数加起来；和 XGBoost，一连串的是否判断，它不怕格子空着。"
         "考法是按时间考，前百分之八十的时间学、后百分之二十考，因为真实场景就是用过去预测未来。"
         "分数上，五类各自算查准、查全和 F1 再取平均。为什么不只看准确率？一个永远说正常的模型准确率有百分之九十一点五，"
         "但一个告警都没找到，所以准确率在这份数据上没有意义。结果：宏平均 F1 从零分提到零点六、再到零点七八。"
         "右图是五类的混淆矩阵，一到三号考得好，四号零点六六、五号零点四四，是两个难点，下一页讲原因。")

# ---------------- 第 8 页 难点与局限 ----------------
s = prs.slides.add_slide(BLANK)
title_bar(s, "分析中解决的问题与现有局限", 8)
tbl = s.shapes.add_table(5, 3, Inches(0.6), Inches(1.4), Inches(12.1), Inches(4.4)).table
tbl.columns[0].width = Inches(4.3)
tbl.columns[1].width = Inches(4.0)
rows = [
    ("问题", "处理办法", "结果"),
    ("数据稀疏（车速仅 6.4% 格子有值）", "统计只基于有值的格子；对比只用样本够的字段", "成因层只有电池电压可检验；转速水温类关系无法检验"),
    ("多类告警同时出现（仅 14 格单一告警）", "对每类分别判定，不强行归并", "保留了全部混合信息，任务定义不需要人为规则"),
    ("按时间考分数低（code5 验证 0.97、测试 0.44）", "诊断为时间分布漂移，非模型问题", "调参、加权、换结构均无效；改细窗口有效，列为后续工作"),
    ("五类告警无故障名称", "成因只写到统计关联为止", "具体故障判断需数据提供方确认"),
]
for i, row in enumerate(rows):
    for j, v in enumerate(row):
        tbl.cell(i, j).text = v
        for p in tbl.cell(i, j).text_frame.paragraphs:
            for run in p.runs:
                run.font.size = Pt(13); run.font.name = "Microsoft YaHei"
notes(s, "这一页说困难和局限。第一，数据稀疏，车速只有百分之六点四的格子有值，所以成因对比里只有电池电压样本够、"
         "能检验，转速水温这些字段和告警的关系不是没关系，是数据不够、检验不了。第二，多类告警同时出现是这份数据的特点，"
         "我们对每类分别判定，没有强行规定一格算哪类。第三，按时间考的时候分数会掉，比如五号验证时零点九七、测试只有零点四四，"
         "我们诊断出原因是规律随时间变了而不是模型不行——调参、加权、换任务结构都试过，分数不动；"
         "把窗口从十秒改细成五秒有效，但清洗和全部分析要跟着重跑，列为后续工作。最后，五类编号没有故障名称，"
         "成因判断只能停在统计关联这一层。")

# ---------------- 第 9 页 总结 ----------------
s = prs.slides.add_slide(BLANK)
title_bar(s, "结论与后续工作", 9)
bullets(s, [
    "数据处理：长表规整为 35,306 × 69 的宽表，逐格核对通过，缺失如实保留",
    "车辆特征：三台车明显不同（快/慢/少动），数据内部自洽（车速转速相关 0.77）",
    "告警：1–4 号“车开着、电压偏高”时成组出现；5 号独立，与点火和电压都无关",
    "模型：1–3 号可自动识别（F1 0.80 以上）；4 号难在部分告警格子与正常格子难以区分，"
    "5 号难在出现规律随时间变化",
], width=Inches(12.3))
bullets(s, [
    "后续：① 以 5 秒窗口重跑全流程（已验证 F1 可提升 0.04）；② 向数据提供方索取告警编码含义；"
    "③ 用长表做告警前后的时序分析，区分“电压先高”还是“告警先报”",
], top=Inches(5.3), width=Inches(12.3), size=16)
notes(s, "总结四句话。数据上，十七万行长表规整成了三万五千行六十九列的宽表，质量逐格核对通过。"
         "车辆上，三台车开得明显不一样，数据本身自洽。告警上，一到四号是车开着、电压偏高时成组出现的一组，"
         "五号独立。模型上，一到三号可以自动识别，四号的困难是部分告警格子和正常格子难以区分，"
         "五号的困难是它的出现规律随时间变化。后续三件事：把窗口改细重跑全流程，我们已经验证这样能提高四分；"
         "向数据提供方要告警编码的含义；再用原始的长表数据做告警前后的时序分析，看电压升高是在告警之前还是之后。"
         "我的汇报到这里，谢谢各位老师。")

out = ROOT / "车辆遥测数据分析汇报.pptx"
prs.save(out)
print(f"已生成：{out}（{len(prs.slides.__iter__.__self__._sldIdLst)} 页）")
